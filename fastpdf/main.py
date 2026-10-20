import os
import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape, portrait
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

def draw_banner_page(c, tid_str, team_str, width, height):
    """Helper function to draw the border and centered content on a canvas page."""
    margin = 36
    c.setLineWidth(2)
    c.setStrokeColorRGB(0, 0, 0)
    c.rect(margin, margin, width - (2 * margin), height - (2 * margin))
    
    center_x = width / 2.0
    center_y = height / 2.0
    
    c.setFont("Helvetica-Bold", 108)
    c.drawCentredString(center_x, center_y + 15, tid_str)
    
    c.setFont("Helvetica", 46)
    c.drawCentredString(center_x, center_y - 25, team_str)

def create_single_pdf(tid, team_name, output_filepath, page_size=landscape(letter)):
    """Generates a single LANDSCAPE PDF with a black border and centered content."""
    c = canvas.Canvas(output_filepath, pagesize=page_size)
    width, height = page_size
    
    tid_text = str(tid) if tid else ""
    team_text = str(team_name) if team_name else ""
    
    draw_banner_page(c, tid_text, team_text, width, height)
    
    c.showPage()
    c.save()

def create_multi_page_pdf(records, output_filepath, page_size=landscape(letter)):
    """Generates a single multi-page PDF containing all banner entries."""
    c = canvas.Canvas(output_filepath, pagesize=page_size)
    width, height = page_size
    
    for record in records:
        tid_str, team_str = record[0], record[1]
        draw_banner_page(c, tid_str, team_str, width, height)
        c.showPage()
        
    c.save()

def create_attendance_table_pdf(records, output_filepath, page_size=portrait(letter)):
    """Generates an uncolored attendance PDF table with left-aligned header."""
    doc = SimpleDocTemplate(
        output_filepath,
        pagesize=page_size,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    header_style = ParagraphStyle(
        'HeaderStyle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        alignment=0,
        spaceAfter=15
    )
    
    th_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=12,
        textColor=colors.black,
        alignment=0
    )
    
    tb_style = ParagraphStyle(
        'TableBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=11,
        alignment=0
    )

    story = []
    
    title = Paragraph("SIH Internal Hackathon Attendance", header_style)
    story.append(title)
    story.append(Spacer(1, 10))

    table_data = [[
        Paragraph("TID", th_style),
        Paragraph("Team Name", th_style),
        Paragraph("Members", th_style),
        Paragraph("Total Present", th_style)
    ]]

    for record in records:
        tid_str, team_str, members_str = record[0], record[1], record[2]
        table_data.append([
            Paragraph(tid_str, tb_style),
            Paragraph(team_str, tb_style),
            Paragraph(members_str, tb_style),
            Paragraph("", tb_style)
        ])

    col_widths = [70, 150, 220, 100]
    
    attendance_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    attendance_table.setStyle(TableStyle([
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('LINEBELOW', (0, 0), (-1, 0), 1.5, colors.black)
    ]))

    story.append(attendance_table)
    doc.build(story)

def create_problem_statements_pdf(records, output_filepath, page_size=portrait(letter)):
    """Generates a black and white table PDF containing TID, Team Name, PS Number, Problem Statement, and Solution."""
    doc = SimpleDocTemplate(
        output_filepath,
        pagesize=page_size,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    header_style = ParagraphStyle(
        'HeaderStylePS',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        alignment=0,
        spaceAfter=15
    )
    
    th_style = ParagraphStyle(
        'TableHeaderPS',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=12,
        textColor=colors.black,
        alignment=0
    )
    
    tb_style = ParagraphStyle(
        'TableBodyPS',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=0
    )

    story = []
    
    title = Paragraph("SIH Problem Statements & Solutions", header_style)
    story.append(title)
    story.append(Spacer(1, 10))

    table_data = [[
        Paragraph("TID", th_style),
        Paragraph("Team Name", th_style),
        Paragraph("PS Number", th_style),
        Paragraph("Problem Statement", th_style),
        Paragraph("Proposed Solution", th_style)
    ]]

    for record in records:
        tid_str, team_str, _, ps_num_str, ps_desc_str, ps_sol_str = record
        table_data.append([
            Paragraph(tid_str, tb_style),
            Paragraph(team_str, tb_style),
            Paragraph(ps_num_str, tb_style),
            Paragraph(ps_desc_str, tb_style),
            Paragraph(ps_sol_str, tb_style)
        ])

    col_widths = [50, 100, 60, 165, 165]
    
    ps_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    
    ps_table.setStyle(TableStyle([
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('LINEBELOW', (0, 0), (-1, 0), 1.5, colors.black)
    ]))

    story.append(ps_table)
    doc.build(story)

def create_team_problem_details_pdf(records, output_filepath, page_size=landscape(letter)):
    """Generates a landscape black and white PDF table with TID, Team Name, Members, PS Number, and Problem Statement."""
    doc = SimpleDocTemplate(
        output_filepath,
        pagesize=page_size,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    header_style = ParagraphStyle(
        'HeaderStyleTeamPS',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        alignment=0,
        spaceAfter=15
    )
    
    th_style = ParagraphStyle(
        'TableHeaderTeamPS',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=12,
        textColor=colors.black,
        alignment=0
    )
    
    tb_style = ParagraphStyle(
        'TableBodyTeamPS',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        alignment=0
    )

    story = []
    
    title = Paragraph("SIH Internal Hackathon Participant List", header_style)
    story.append(title)
    story.append(Spacer(1, 10))

    table_data = [[
        Paragraph("TID", th_style),
        Paragraph("Team Name", th_style),
        Paragraph("Members", th_style),
        Paragraph("PS Number", th_style),
        Paragraph("Problem Statement", th_style)
    ]]

    for record in records:
        tid_str, team_str, members_str, ps_num_str, ps_desc_str, _ = record
        table_data.append([
            Paragraph(tid_str, tb_style),
            Paragraph(team_str, tb_style),
            Paragraph(members_str, tb_style),
            Paragraph(ps_num_str, tb_style),
            Paragraph(ps_desc_str, tb_style)
        ])

    # Printable landscape width: 792 - 72 = 720pt total
    col_widths = [60, 130, 190, 70, 270]
    
    team_ps_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    team_ps_table.setStyle(TableStyle([
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('LINEBELOW', (0, 0), (-1, 0), 1.5, colors.black)
    ]))

    story.append(team_ps_table)
    doc.build(story)

def create_teams_txt_file(records, output_filepath):
    """Generates a text file listing TID and Team Name side-by-side."""
    with open(output_filepath, "w", encoding="utf-8") as f:
        f.write(f"{'TID':<15}\t{'Team Name'}\n")
        f.write("-" * 50 + "\n")
        for record in records:
            f.write(f"{record[0]:<15}\t{record[1]}\n")

def batch_generate_pdfs_from_excel(excel_filename, output_folder="generated_pdfs"):
    if not os.path.exists(excel_filename):
        print(f"Error: File '{excel_filename}' not found.")
        return

    if not os.path.exists(output_folder):
        os.makedirs(output_folder)
        
    wb = openpyxl.load_workbook(excel_filename, data_only=True)
    sheet = wb.active
    
    headers = [str(cell.value).strip().lower() if cell.value else "" for cell in sheet[1]]
    
    tid_col_idx = None
    team_col_idx = None
    members_col_idx = None
    ps_num_col_idx = None
    ps_desc_col_idx = None
    ps_sol_col_idx = None
    
    for idx, header in enumerate(headers):
        if header in ['tid', 'team id', 'id']:
            tid_col_idx = idx + 1
        elif header in ['team name', 'team', 'team_name']:
            team_col_idx = idx + 1
        elif header in ['members', 'team members', 'member_names', 'member']:
            members_col_idx = idx + 1
        elif header in ['ps number', 'ps_number', 'ps num', 'ps_id', 'ps id']:
            ps_num_col_idx = idx + 1
        elif header in ['problem statement', 'problem_statement', 'ps', 'ps description', 'statement']:
            ps_desc_col_idx = idx + 1
        elif header in ['proposed solution', 'proposed_solution', 'solution', 'sol', 'solution description']:
            ps_sol_col_idx = idx + 1
            
    if not tid_col_idx or not team_col_idx:
        print("Error: Could not locate required TID/Team Name columns.")
        return

    records = []
    count = 0
    
    for row in range(2, sheet.max_row + 1):
        tid_val = sheet.cell(row=row, column=tid_col_idx).value
        team_val = sheet.cell(row=row, column=team_col_idx).value
        members_val = sheet.cell(row=row, column=members_col_idx).value if members_col_idx else ""
        ps_num_val = sheet.cell(row=row, column=ps_num_col_idx).value if ps_num_col_idx else ""
        ps_desc_val = sheet.cell(row=row, column=ps_desc_col_idx).value if ps_desc_col_idx else ""
        ps_sol_val = sheet.cell(row=row, column=ps_sol_col_idx).value if ps_sol_col_idx else ""
        
        if tid_val is None and team_val is None:
            continue
            
        tid_str = str(tid_val).strip() if tid_val is not None else ""
        team_str = str(team_val).strip() if team_val is not None else ""
        members_str = str(members_val).strip() if members_val is not None else ""
        ps_num_str = str(ps_num_val).strip() if ps_num_val is not None else ""
        ps_desc_str = str(ps_desc_val).strip() if ps_desc_val is not None else ""
        ps_sol_str = str(ps_sol_val).strip() if ps_sol_val is not None else ""
        
        records.append((tid_str, team_str, members_str, ps_num_str, ps_desc_str, ps_sol_str))
        
        clean_tid = "".join([c for c in tid_str if c.isalnum() or c in ('_', '-')]).rstrip()
        filename = f"{clean_tid if clean_tid else f'Row_{row}'}.pdf"
        filepath = os.path.join(output_folder, filename)
            
        create_single_pdf(tid_str, team_str, filepath)
        count += 1
        print(f"[{count}] Generated: {filepath}")

    if records:
        # Combined Banners PDF
        combined_pdf_path = os.path.join(output_folder, "LIRH banner.pdf")
        create_multi_page_pdf(records, combined_pdf_path)
        print(f"\nGenerated combined banner PDF: {combined_pdf_path}")
        
        # Attendance Table PDF
        attendance_pdf_path = os.path.join(output_folder, "Attendance_Sheet.pdf")
        create_attendance_table_pdf(records, attendance_pdf_path)
        print(f"Generated attendance PDF: {attendance_pdf_path}")
        
        # Problem Statements & Solutions PDF
        ps_pdf_path = os.path.join(output_folder, "Problem_Statements_Sheet.pdf")
        create_problem_statements_pdf(records, ps_pdf_path)
        print(f"Generated problem statements PDF: {ps_pdf_path}")

        # Team Problem Details PDF (TID, Team Name, Members, PS Number, Problem Statement)
        team_ps_pdf_path = os.path.join(output_folder, "Team_Problem_Statements_Sheet.pdf")
        create_team_problem_details_pdf(records, team_ps_pdf_path)
        print(f"Generated team problem statement details PDF: {team_ps_pdf_path}")
        
        # Text File output
        teams_txt_path = os.path.join(output_folder, "Teams_List.txt")
        create_teams_txt_file(records, teams_txt_path)
        print(f"Generated teams text file: {teams_txt_path}")

    print(f"\nSuccess! Total {count} individual files and summary documents generated in '{output_folder}'.")

if __name__ == "__main__":
    try:
        filename = input('Enter "Path to .xlsx" file : ').strip("'\"")
        batch_generate_pdfs_from_excel(filename)
    except FileNotFoundError:
        print("invalid path provided")
    except Exception as e:
        print(f"Error receiving inputs: {e}")